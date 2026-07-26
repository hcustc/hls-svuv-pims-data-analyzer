from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from bl03u_masstool.core.export_pie_results import export_pie_results_to_excel


def test_export_pie_results_keeps_missing_ie_blank(tmp_path):
    dialog = SimpleNamespace(
        all_fit_results={
            30: {
                "success": True,
                "r_squared": 0.9,
                "model": {
                    "species": [
                        {
                            "species": "Unknown isomer",
                            "ie": None,
                            "coefficient": 1.0,
                            "contribution_percent": 100.0,
                        }
                    ]
                },
            }
        }
    )
    output_path = tmp_path / "pie_results.xlsx"

    result = export_pie_results_to_excel(dialog, output_path)

    assert result["success"] is True
    exported = pd.read_excel(output_path)
    assert pd.isna(exported.loc[0, "电离能(eV)"])


def test_export_pie_results_falls_back_to_ionization_energy_alias(tmp_path):
    dialog = SimpleNamespace(
        all_fit_results={
            44: {
                "success": True,
                "r_squared": 0.95,
                "model": {
                    "species": [
                        {
                            "species": "CO2",
                            "ie": None,
                            "ionization_energy": 13.777,
                            "coefficient": 2.0,
                            "contribution_percent": 100.0,
                        }
                    ]
                },
            }
        }
    )
    output_path = tmp_path / "pie_results.xlsx"

    result = export_pie_results_to_excel(dialog, output_path)

    assert result["success"] is True
    exported = pd.read_excel(output_path)
    assert exported.loc[0, "电离能(eV)"] == 13.777


def test_export_pie_results_adds_structure_oriented_species_sheet(tmp_path):
    dialog = SimpleNamespace(
        database=[
            {
                "id": 7,
                "mz": 28,
                "species": "Ethylene",
                "formula": "C2H4",
                "smiles": "C=C",
            }
        ],
        all_fit_results={
            28: {
                "success": True,
                "r_squared": 0.99,
                "model": {
                    "species": [
                        {
                            "id": 7,
                            "species": "Ethylene",
                            "ie": 10.51,
                            "coefficient": 1.2,
                            "contribution_percent": 100.0,
                        }
                    ]
                },
            }
        },
    )
    output_path = tmp_path / "pie_results.xlsx"

    result = export_pie_results_to_excel(dialog, output_path)

    assert result["success"] is True
    assert result["species_count"] == 1
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["Sheet1", "物种表"]
    species = pd.read_excel(output_path, sheet_name="物种表")
    assert species.columns.tolist() == [
        "精确m/z",
        "m/z",
        "Formula",
        "Name",
        "SMILES",
        "Structure",
    ]
    assert species.loc[0, "m/z"] == 28
    assert species.loc[0, "Formula"] == "C2H4"
    assert species.loc[0, "Name"] == "Ethylene"
    assert species.loc[0, "SMILES"] == "C=C"
    assert len(workbook["物种表"]._images) == 1
    assert workbook["物种表"].row_dimensions[2].height == 82


def test_export_species_sheet_keeps_invalid_smiles_without_image(tmp_path):
    dialog = SimpleNamespace(
        database=[],
        all_fit_results={
            15: {
                "success": True,
                "r_squared": 0.8,
                "model": {
                    "species": [
                        {
                            "species": "Methyl radical",
                            "formula": "CH3",
                            "smiles": "not-a-smiles",
                            "coefficient": 1.0,
                            "contribution_percent": 100.0,
                        }
                    ]
                },
            }
        },
    )
    output_path = tmp_path / "pie_results.xlsx"

    result = export_pie_results_to_excel(dialog, output_path)

    assert result["success"] is True
    species = pd.read_excel(output_path, sheet_name="物种表")
    assert species.loc[0, "SMILES"] == "not-a-smiles"
    assert len(load_workbook(output_path)["物种表"]._images) == 0


def test_export_keeps_two_precise_peaks_with_same_nominal_mass(tmp_path):
    exact_mz_values = [227.58715739409433, 228.02311680380544]
    dialog = SimpleNamespace(
        database=[],
        curves={
            exact_mz: {
                "mz": 228,
                "mz_rounded": 228,
                "mz_exact_mean": exact_mz,
                "has_nominal_collision": True,
            }
            for exact_mz in exact_mz_values
        },
        all_fit_results={
            exact_mz: {
                "success": True,
                "r_squared": 0.9,
                "model": {
                    "species": [
                        {
                            "species": f"candidate-{index}",
                            "coefficient": 1.0,
                            "contribution_percent": 100.0,
                        }
                    ]
                },
            }
            for index, exact_mz in enumerate(exact_mz_values)
        },
    )
    output_path = tmp_path / "precise_pie_results.xlsx"

    result = export_pie_results_to_excel(dialog, output_path)

    assert result["success"] is True
    exported = pd.read_excel(output_path)
    assert exported["质量数"].tolist() == [228, 228]
    assert exported["精确m/z"].tolist() == pytest.approx(exact_mz_values)
