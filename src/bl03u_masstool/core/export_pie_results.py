from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from bl03u_masstool.core.pie_analysis import species_ionization_energy_value


SUMMARY_COLUMNS = [
    "精确m/z",
    "质量数",
    "物种名称",
    "电离能(eV)",
    "匹配系数",
    "贡献比例(%)",
    "R²",
]
SPECIES_COLUMNS = ["精确m/z", "m/z", "Formula", "Name", "SMILES", "Structure"]


def _species_database_metadata(pie_dialog, mz: int, species_info: dict) -> tuple[str, str]:
    formula = str(species_info.get("formula") or "").strip()
    smiles = str(species_info.get("smiles") or "").strip()
    if formula and smiles:
        return formula, smiles

    species_name = str(species_info.get("species") or "").strip().casefold()
    species_ids = {
        int(value)
        for value in (
            [species_info.get("id")]
            + list(species_info.get("ids") or [])
        )
        if value is not None and str(value).lstrip("-").isdigit()
    }
    for candidate in getattr(pie_dialog, "database", []) or []:
        try:
            candidate_mz = int(candidate.get("mz"))
        except (TypeError, ValueError):
            continue
        candidate_id = candidate.get("id")
        try:
            id_matches = candidate_id is not None and int(candidate_id) in species_ids
        except (TypeError, ValueError):
            id_matches = False
        name_matches = (
            species_name
            and str(candidate.get("species") or "").strip().casefold() == species_name
        )
        if candidate_mz == int(mz) and (id_matches or name_matches):
            formula = formula or str(candidate.get("formula") or "").strip()
            smiles = smiles or str(candidate.get("smiles") or "").strip()
            if formula and smiles:
                break
    return formula, smiles


def _collect_identified_species(pie_dialog) -> list[dict]:
    records: list[dict] = []
    curves = getattr(pie_dialog, "curves", {}) or {}
    for curve_key, fit_result in (getattr(pie_dialog, "all_fit_results", {}) or {}).items():
        if not fit_result.get("success"):
            continue
        model = fit_result.get("model")
        if not model:
            continue
        curve = curves.get(curve_key, {})
        nominal_mz = int(
            round(float(curve.get("mz_rounded", curve.get("mz", curve_key))))
        )
        exact_mz = float(curve.get("mz_exact_mean", curve_key))
        species_list = model.get("species", [])
        r_squared = fit_result.get("r_squared", model.get("r_squared", 0))
        if not species_list:
            records.append(
                {
                    "精确m/z": exact_mz,
                    "质量数": nominal_mz,
                    "物种名称": "Unknown",
                    "电离能(eV)": None,
                    "匹配系数": 0.0,
                    "贡献比例(%)": 0.0,
                    "R²": r_squared,
                    "Formula": "",
                    "SMILES": "",
                }
            )
            continue
        for species_info in species_list:
            formula, smiles = _species_database_metadata(
                pie_dialog,
                nominal_mz,
                species_info,
            )
            records.append(
                {
                    "精确m/z": exact_mz,
                    "质量数": nominal_mz,
                    "物种名称": species_info.get("species", "Unknown"),
                    "电离能(eV)": species_ionization_energy_value(species_info),
                    "匹配系数": species_info.get("coefficient", 0),
                    "贡献比例(%)": species_info.get("contribution_percent", 0),
                    "R²": r_squared,
                    "Formula": formula,
                    "SMILES": smiles,
                }
            )
    return records


def _structure_png(smiles: str) -> BytesIO | None:
    if not smiles:
        return None
    try:
        from rdkit import Chem, rdBase
        from rdkit.Chem import Draw, rdDepictor

        with rdBase.BlockLogs():
            molecule = Chem.MolFromSmiles(smiles)
        if molecule is None:
            return None
        rdDepictor.Compute2DCoords(molecule)
        image = Draw.MolToImage(molecule, size=(220, 130), fitImage=True)
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        return buffer
    except (ImportError, ValueError, RuntimeError):
        return None


def _format_summary_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="E9EFF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)
    exact_mz_column = next(
        (cell.column for cell in worksheet[1] if cell.value == "精确m/z"),
        None,
    )
    if exact_mz_column is not None:
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=exact_mz_column).number_format = "0.000000000000"


def _write_species_sheet(writer, species_df: pd.DataFrame) -> None:
    species_df.to_excel(writer, index=False, sheet_name="物种表")
    worksheet = writer.sheets["物种表"]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    header_fill = PatternFill("solid", fgColor="E9EFF7")
    separator = Side(style="thin", color="AAB7C4")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=separator)
    worksheet.row_dimensions[1].height = 24
    widths = {"A": 19, "B": 10, "C": 18, "D": 32, "E": 42, "F": 31}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    image_buffers: list[BytesIO] = []
    previous_mz = None
    for row_number, record in enumerate(species_df.to_dict("records"), start=2):
        current_mz = record["精确m/z"]
        top_border = Border(top=separator) if previous_mz is not None and current_mz != previous_mz else Border()
        for column in range(1, 7):
            cell = worksheet.cell(row=row_number, column=column)
            cell.alignment = Alignment(
                horizontal="center" if column in {1, 2, 3, 6} else "left",
                vertical="center",
                wrap_text=column in {4, 5},
            )
            cell.border = top_border
        worksheet.cell(row=row_number, column=1).number_format = "0.000000000000"
        worksheet.row_dimensions[row_number].height = 82
        structure = _structure_png(str(record.get("SMILES") or ""))
        if structure is not None:
            image_buffers.append(structure)
            image = WorksheetImage(structure)
            image.width = 172
            image.height = 102
            image.anchor = f"F{row_number}"
            worksheet.add_image(image)
        previous_mz = current_mz
    # WorksheetImage keeps references to these streams until the workbook save.


def export_pie_results_to_excel(pie_dialog, output_path: str | Path) -> dict:
    """Export the numerical PIE result and a concise structure-oriented species table."""
    output_path = Path(output_path)
    records = _collect_identified_species(pie_dialog)
    if not records:
        return {"success": False, "message": "没有找到拟合结果数据", "count": 0}

    identified_df = pd.DataFrame(records).sort_values(
        ["精确m/z", "贡献比例(%)"], ascending=[True, False]
    )
    summary_df = identified_df.drop_duplicates(
        subset=["精确m/z", "物种名称"], keep="first"
    )

    export_df = summary_df[SUMMARY_COLUMNS].copy()
    export_df["电离能(eV)"] = pd.to_numeric(export_df["电离能(eV)"], errors="coerce").round(4)
    export_df["匹配系数"] = pd.to_numeric(export_df["匹配系数"], errors="coerce").round(6)
    export_df["贡献比例(%)"] = pd.to_numeric(export_df["贡献比例(%)"], errors="coerce").round(2)
    export_df["R²"] = pd.to_numeric(export_df["R²"], errors="coerce").round(6)
    export_df["精确m/z"] = pd.to_numeric(export_df["精确m/z"], errors="coerce")
    export_df["质量数"] = export_df["质量数"].astype(int)

    species_df = identified_df.rename(columns={"质量数": "m/z", "物种名称": "Name"})[
        SPECIES_COLUMNS[:-1]
    ].copy()
    species_df["Structure"] = ""
    species_df = species_df.drop_duplicates(
        subset=["精确m/z", "Formula", "Name", "SMILES"], keep="first"
    ).sort_values(["精确m/z", "Name"], kind="stable")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="Sheet1")
        _format_summary_sheet(writer.sheets["Sheet1"])
        _write_species_sheet(writer, species_df)

    return {
        "success": True,
        "message": (
            f"成功导出 {len(export_df)} 条鉴定记录和 "
            f"{len(species_df)} 条物种记录到 {output_path}"
        ),
        "count": len(export_df),
        "species_count": len(species_df),
        "file_path": str(output_path),
    }
