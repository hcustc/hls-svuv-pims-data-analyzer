from __future__ import annotations

from pathlib import Path
from typing import Iterable

import numpy as np
import openpyxl
import pandas as pd

try:
    _np_trapezoid = np.trapezoid
except AttributeError:
    _np_trapezoid = np.trapz

from .config import load_peak_integration_config
from .peak_detection import Peak, fit_gaussian
from .spectrum_io import extract_header_numbers, list_spectrum_files, read_spectrum


_INTEGRATION_METHODS = {"sum_counts", "baseline", "gaussian"}


def resolve_integration_method(
    integration_method: str | None,
    *,
    prefer_gaussian: bool = True,
) -> str:
    """Resolve a method while giving an explicit valid selection precedence."""
    explicit = str(integration_method or "").strip().lower()
    if explicit in _INTEGRATION_METHODS:
        return explicit
    return "gaussian" if prefer_gaussian else "sum_counts"


def baseline_corrected_area(y_data: Iterable[float], left: int, right: int) -> float:
    data = np.asarray(y_data, dtype=float)
    left = max(0, int(left))
    right = min(len(data) - 1, int(right))
    if left > right or data.size == 0:
        return 0.0
    segment = data[left : right + 1]
    if segment.size == 0:
        return 0.0
    return float(_np_trapezoid(segment - np.min(segment)))


def summed_counts_area(y_data: Iterable[float], left: int, right: int) -> float:
    data = np.asarray(y_data, dtype=float)
    left = max(0, int(left))
    right = min(len(data) - 1, int(right))
    if left > right or data.size == 0:
        return 0.0
    return float(np.sum(data[left : right + 1]))


def gaussian_area(amplitude: float, fwhm: float) -> float:
    return float(max(0.0, amplitude * fwhm * np.sqrt(np.pi / (4 * np.log(2)))))


def integrate_peak(y_data: Iterable[float], peak: Peak, *, prefer_gaussian: bool = True) -> float:
    area, _method = integrate_peak_with_method(y_data, peak, prefer_gaussian=prefer_gaussian)
    return area


def integrate_peak_with_method(
    y_data: Iterable[float],
    peak: Peak,
    *,
    prefer_gaussian: bool = True,
    integration_method: str | None = None,
) -> tuple[float, str]:
    method = resolve_integration_method(
        integration_method,
        prefer_gaussian=prefer_gaussian,
    )
    if method == "gaussian":
        window_size = max(5, min(30, int(peak.right_bound) - int(peak.left_bound) + 5))
        fit = fit_gaussian(y_data, int(round(peak.index)), window_size)
        if fit is not None:
            area = gaussian_area(fit.amplitude, fit.fwhm)
            if area > 0:
                return area, "gaussian"
        method = "sum_counts"
    if method == "baseline":
        return baseline_corrected_area(y_data, peak.left_bound, peak.right_bound), "baseline"
    return summed_counts_area(y_data, peak.left_bound, peak.right_bound), "sum_counts"


def integrate_peaks_with_method(
    y_data: Iterable[float],
    peaks: Iterable[Peak],
    *,
    prefer_gaussian: bool = True,
    integration_method: str | None = None,
) -> list[tuple[float, str]]:
    """Integrate one spectrum against many peak windows with stable semantics.

    Sum and baseline modes convert the spectrum to an array once.  Temperature
    scans commonly integrate hundreds of reference peaks for every spectrum,
    so avoiding repeated dispatch and bounds normalization reduces overhead
    without changing the NumPy operations or their order.
    """
    data = np.asarray(y_data, dtype=float)
    peak_list = list(peaks)
    method = resolve_integration_method(
        integration_method,
        prefer_gaussian=prefer_gaussian,
    )
    if method == "gaussian":
        return [
            integrate_peak_with_method(
                data,
                peak,
                prefer_gaussian=True,
                integration_method="gaussian",
            )
            for peak in peak_list
        ]

    results: list[tuple[float, str]] = []
    data_size = int(data.size)
    for peak in peak_list:
        left = max(0, int(peak.left_bound))
        right = min(data_size - 1, int(peak.right_bound))
        if left > right or data_size == 0:
            results.append((0.0, method))
            continue
        segment = data[left : right + 1]
        if method == "baseline":
            area = 0.0 if segment.size == 0 else float(_np_trapezoid(segment - np.min(segment)))
        else:
            area = float(np.sum(segment))
        results.append((area, method))
    return results


def integrate_peaks(y_data: Iterable[float], peaks: Iterable[Peak], *, prefer_gaussian: bool = True) -> list[dict]:
    return [
        {
            "mz": peak.mz,
            "time": peak.time,
            "left_bound": peak.left_bound,
            "right_bound": peak.right_bound,
            "area": integrate_peak(y_data, peak, prefer_gaussian=prefer_gaussian),
        }
        for peak in peaks
    ]


def load_peak_config_excel(path: str | Path, sheet_name: str = "卡峰配置") -> dict:
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[sheet_name]
    config: dict = {}
    for row in range(2, sheet.max_row + 1):
        species = sheet.cell(column=1, row=row).value
        if species is None:
            continue
        config.setdefault(species, {})
        for col in range(3, 6):
            key = sheet.cell(column=col, row=1).value
            config[species][key] = sheet.cell(column=col, row=row).value
    return config


def load_peak_config_yaml(path: str | Path | None = None) -> dict:
    config: dict = {}
    for peak in load_peak_integration_config(path):
        species = peak["formula"]
        config[species] = {
            "M/Z": peak["mz"],
            "Peak": peak["peak"],
            "Start": peak["start"],
            "End": peak["end"],
        }
    return config


def load_peak_config(path: str | Path | None = None) -> dict:
    if path is None:
        return load_peak_config_yaml()
    path = Path(path)
    if path.suffix.lower() in {".yaml", ".yml"}:
        return load_peak_config_yaml(path)
    return load_peak_config_excel(path)


def integrate_by_config(y_data: Iterable[float], config: dict, *, mode: str = "area") -> dict:
    data = np.asarray(list(y_data), dtype=float)
    result = {}
    for species, info in config.items():
        start = int(info["Start"])
        end = int(info["End"])
        segment = data[start : end + 1]
        if segment.size == 0:
            result[species] = 0.0
        elif mode == "height":
            result[species] = float(np.max(segment))
        else:
            result[species] = float(_np_trapezoid(segment, dx=0.01))
    return result


def integrate_folder_by_config(
    folder: str | Path,
    config: dict,
    *,
    suffixes: tuple[str, ...] = (".txt",),
    normalize_header_index: int | None = 1,
) -> pd.DataFrame:
    rows = []
    for path in list_spectrum_files(folder, suffixes):
        spectrum = read_spectrum(path, header_lines=10, trim_start=0)
        values = integrate_by_config(spectrum.y, config)
        if normalize_header_index is not None:
            header_numbers = extract_header_numbers(path)
            if len(header_numbers) > normalize_header_index and header_numbers[normalize_header_index] != 0:
                values = {key: value / header_numbers[normalize_header_index] for key, value in values.items()}
        rows.append({"file": path.name, **values})
    return pd.DataFrame(rows)
